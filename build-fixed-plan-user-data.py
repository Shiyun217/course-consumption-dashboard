import json
import os
from collections import Counter, defaultdict
from datetime import date, datetime

from openpyxl import load_workbook


BIND_PATH = os.environ["FP_BIND"]
PAY_PATH = os.environ["FP_PAY"]
OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "fixed-plan-user-data.json")
JS_OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "fixed-plan-user-data.js")
START_MONTH = "2025-02"
END_MONTH = "2026-07"
PORTS = ("CC", "SS", "LP", "student", "other")


def normalize_id(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def to_datetime(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, str) and value.strip():
        text = value.strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                pass
    return None


def month_key(value):
    parsed = to_datetime(value)
    return parsed.strftime("%Y-%m") if parsed else None


def month_index(value):
    year, month = map(int, value.split("-"))
    return year * 12 + month


def month_range(start, end):
    current = month_index(start)
    final = month_index(end)
    result = []
    while current <= final:
        year, offset = divmod(current - 1, 12)
        result.append(f"{year:04d}-{offset + 1:02d}")
        current += 1
    return result


def normalize_port(value):
    text = str(value or "").strip()
    lowered = text.lower()
    if lowered == "cc":
        return "CC"
    if lowered == "ss":
        return "SS"
    if lowered == "lp":
        return "LP"
    if lowered == "student":
        return "student"
    return "other"


def load_pay_students():
    workbook = load_workbook(PAY_PATH, read_only=True, data_only=True)
    sheet = workbook.active
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    positions = {name: index for index, name in enumerate(headers)}
    id_index = positions["stdt_id"]
    pay_index = positions["first_1v1_nml_pay_date"]
    students = {}
    duplicates = 0
    conflicting_months = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        student_id = normalize_id(row[id_index])
        pay_date = to_datetime(row[pay_index])
        if not student_id or not pay_date:
            continue
        existing = students.get(student_id)
        if existing is not None:
            duplicates += 1
            if month_key(existing) != month_key(pay_date):
                conflicting_months += 1
            if pay_date < existing:
                students[student_id] = pay_date
        else:
            students[student_id] = pay_date
    return students, {"duplicate_pay_rows": duplicates, "conflicting_pay_months": conflicting_months}


def load_first_bindings():
    workbook = load_workbook(BIND_PATH, read_only=True, data_only=True)
    sheet = workbook.active
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    positions = {name: index for index, name in enumerate(headers)}
    id_index = positions["b.student_id"]
    bind_index = positions["固定绑定时间"]
    pay_index = positions["首单时间"]
    port_index = positions["bind_operator_type"]
    first_bindings = {}
    bind_pay_months = {}
    raw_ports = Counter()
    valid_rows = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        student_id = normalize_id(row[id_index])
        bind_time = to_datetime(row[bind_index])
        if not student_id or not bind_time:
            continue
        valid_rows += 1
        raw_ports[str(row[port_index] or "").strip()] += 1
        candidate = {
            "time": bind_time,
            "port": normalize_port(row[port_index]),
            "raw_port": str(row[port_index] or "").strip(),
        }
        existing = first_bindings.get(student_id)
        if existing is None or bind_time < existing["time"]:
            first_bindings[student_id] = candidate
            bind_pay_months[student_id] = month_key(row[pay_index])
    return first_bindings, bind_pay_months, {
        "valid_binding_rows": valid_rows,
        "unique_bound_students": len(first_bindings),
        "raw_operator_types": dict(raw_ports.most_common()),
    }


def rate(count, denominator):
    return count / denominator if denominator else 0


pay_students, pay_quality = load_pay_students()
first_bindings, bind_pay_months, bind_quality = load_first_bindings()
months = month_range(START_MONTH, END_MONTH)
cohort_students = defaultdict(list)
for student_id, pay_date in pay_students.items():
    cohort = month_key(pay_date)
    if START_MONTH <= cohort <= END_MONTH:
        cohort_students[cohort].append(student_id)

rows = []
pre_pay_count = 0
pay_month_mismatch = 0
unmatched_bound_students = 0
for student_id, binding in first_bindings.items():
    if student_id not in pay_students:
        unmatched_bound_students += 1
        continue
    source_month = month_key(pay_students[student_id])
    embedded_month = bind_pay_months.get(student_id)
    if embedded_month and source_month != embedded_month:
        pay_month_mismatch += 1

for cohort in months:
    students = cohort_students.get(cohort, [])
    denominator = len(students)
    lifecycle_counts = Counter()
    port_counts = {stage: Counter() for stage in ("M0", "M1", "M2plus")}
    for student_id in students:
        binding = first_bindings.get(student_id)
        if not binding:
            continue
        binding_month = month_key(binding["time"])
        delta = month_index(binding_month) - month_index(cohort)
        if delta < 0:
            pre_pay_count += 1
            continue
        stage = "M0" if delta == 0 else "M1" if delta == 1 else "M2plus"
        lifecycle_counts[stage] += 1
        port_counts[stage][binding["port"]] += 1
    total = sum(lifecycle_counts.values())
    stage_payload = {}
    for stage in ("M0", "M1", "M2plus"):
        stage_payload[stage] = {
            "count": lifecycle_counts[stage],
            "rate": rate(lifecycle_counts[stage], denominator),
            "ports": {
                port: {
                    "count": port_counts[stage][port],
                    "rate": rate(port_counts[stage][port], denominator),
                }
                for port in PORTS
            },
        }
    rows.append({
        "cohort": cohort,
        "new_students": denominator,
        "M0": stage_payload["M0"],
        "M1": stage_payload["M1"],
        "M2plus": stage_payload["M2plus"],
        "total": {"count": total, "rate": rate(total, denominator)},
        "maturity": {
            "M0": True,
            "M1": month_index(cohort) + 1 <= month_index(END_MONTH),
            "M2plus": month_index(cohort) + 2 <= month_index(END_MONTH),
        },
    })

aggregate_denominator = sum(row["new_students"] for row in rows)
aggregate = {}
for stage in ("M0", "M1", "M2plus"):
    count = sum(row[stage]["count"] for row in rows)
    aggregate[stage] = {
        "count": count,
        "rate": rate(count, aggregate_denominator),
        "ports": {
            port: {
                "count": sum(row[stage]["ports"][port]["count"] for row in rows),
                "rate": rate(sum(row[stage]["ports"][port]["count"] for row in rows), aggregate_denominator),
            }
            for port in PORTS
        },
    }
aggregate_total = sum(aggregate[stage]["count"] for stage in aggregate)
aggregate["total"] = {"count": aggregate_total, "rate": rate(aggregate_total, aggregate_denominator)}

result = {
    "source_as_of": END_MONTH,
    "range": {"start": START_MONTH, "end": END_MONTH},
    "ports": list(PORTS),
    "rows": rows,
    "aggregate": {"new_students": aggregate_denominator, **aggregate},
    "quality": {
        **pay_quality,
        **bind_quality,
        "pay_students": len(pay_students),
        "cohort_students": aggregate_denominator,
        "pre_pay_first_bindings_excluded": pre_pay_count,
        "binding_pay_month_mismatches": pay_month_mismatch,
        "bound_students_missing_from_pay_table": unmatched_bound_students,
    },
}

with open(OUTPUT_PATH, "w", encoding="utf-8") as handle:
    json.dump(result, handle, ensure_ascii=False, indent=2)

with open(JS_OUTPUT_PATH, "w", encoding="utf-8") as handle:
    handle.write("window.FIXED_PLAN_USER_DATA = ")
    json.dump(result, handle, ensure_ascii=False, separators=(",", ":"))
    handle.write(";\n")

print(json.dumps({
    "output": OUTPUT_PATH,
    "aggregate": result["aggregate"],
    "quality": result["quality"],
    "rows": result["rows"],
}, ensure_ascii=False, indent=2))
