import csv
import calendar
import json
import os
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


BIND_PATH = Path(os.environ["FP_BIND"])
PAY_PATH = Path(os.environ["FP_PAY"])
BASE_DIR = Path(__file__).resolve().parent
JSON_OUTPUT = BASE_DIR / "fixed-plan-m01-all-data.json"
JS_OUTPUT = BASE_DIR / "fixed-plan-m01-all-data.js"
START_MONTH = os.environ.get("FP_START_MONTH", "2026-08")
END_MONTH = os.environ.get("FP_END_MONTH", START_MONTH)
end_year, end_month = map(int, END_MONTH.split("-"))
SOURCE_AS_OF = os.environ.get(
    "FP_SOURCE_AS_OF",
    f"{END_MONTH}-{calendar.monthrange(end_year, end_month)[1]:02d}",
)
TARGET_BASELINE_MONTH = os.environ.get("FP_TARGET_BASELINE_MONTH", "2026-07")
PORTS = ("CC", "SS", "LP", "student", "other")


def normalize_id(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    normalized = str(value).strip()
    return normalized[:-2] if normalized.endswith(".0") else normalized


def to_datetime(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, str) and value.strip() and value.strip() != r"\N":
        cleaned = value.strip().split(".", 1)[0]
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d"):
            try:
                return datetime.strptime(cleaned, fmt)
            except ValueError:
                pass
    return None


def month_key(value):
    parsed = to_datetime(value)
    return parsed.strftime("%Y-%m") if parsed else None


def month_index(value):
    year, month = map(int, value.split("-"))
    return year * 12 + month


def month_from_index(value):
    year, offset = divmod(value - 1, 12)
    return f"{year:04d}-{offset + 1:02d}"


def month_range(start, end):
    return [month_from_index(value) for value in range(month_index(start), month_index(end) + 1)]


def normalize_port(operator_type, operator_group_type=""):
    lowered = str(operator_type or "").strip().lower()
    group_type = str(operator_group_type or "").strip().lower()
    # New exports record CC as operator_type=other and operator_group_type=cc.
    if group_type == "cc":
        return "CC"
    if lowered in ("cc", "ss", "lp"):
        return lowered.upper()
    if lowered == "student":
        return "student"
    return "other"


def rate(count, denominator):
    return count / denominator if denominator else 0


def iter_records(path):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            yield from csv.DictReader(handle)
        return
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    for values in rows:
        yield dict(zip(headers, values))


def first_value(row, *names):
    for name in names:
        value = row.get(name)
        if value not in (None, "", r"\N"):
            return value
    return None


def load_pay_students():
    students = {}
    for row in iter_records(PAY_PATH):
        student_id = normalize_id(first_value(row, "stdt_id", "student_id"))
        pay_date = to_datetime(first_value(row, "first_1v1_nml_pay_date", "首单时间"))
        if not student_id or not pay_date:
            continue
        current = students.get(student_id)
        if current is None or pay_date < current:
            students[student_id] = pay_date
    return students


def load_binding_touches():
    touches_by_month = defaultdict(lambda: defaultdict(set))
    valid_rows = 0
    for row in iter_records(BIND_PATH):
        student_id = normalize_id(first_value(row, "b.student_id", "student_id"))
        binding_month = month_key(first_value(row, "固定绑定时间", "bind_add_time"))
        if not student_id or not binding_month:
            continue
        valid_rows += 1
        port = normalize_port(row.get("bind_operator_type"), row.get("bind_operator_group_type"))
        touches_by_month[binding_month][port].add(student_id)
    return touches_by_month, valid_rows


existing = json.loads(JSON_OUTPUT.read_text(encoding="utf-8-sig")) if JSON_OUTPUT.exists() else {}
pay_students = load_pay_students()
touches_by_month, valid_binding_rows = load_binding_touches()
pay_by_month = defaultdict(set)
for student_id, pay_date in pay_students.items():
    pay_by_month[month_key(pay_date)].add(student_id)

updated_rows = []
for report_month in month_range(START_MONTH, END_MONTH):
    previous_month = month_from_index(month_index(report_month) - 1)
    eligible = pay_by_month[previous_month] | pay_by_month[report_month]
    port_students = {}
    for port in PORTS:
        touched = touches_by_month[previous_month][port] | touches_by_month[report_month][port]
        port_students[port] = eligible & touched
    union_students = set().union(*port_students.values())
    touch_sum = sum(len(values) for values in port_students.values())
    multi_port_students = sum(
        1 for student_id in union_students
        if sum(student_id in port_students[port] for port in PORTS) > 1
    )
    denominator = len(eligible)
    updated_rows.append({
        "month": report_month,
        "new_students": denominator,
        "bound_students": len(union_students),
        "binding_rate": rate(len(union_students), denominator),
        "port_counts": {port: len(port_students[port]) for port in PORTS},
        "rates": {port: rate(len(port_students[port]), denominator) for port in PORTS},
        "touch_sum": touch_sum,
        "touch_sum_rate": rate(touch_sum, denominator),
        "overlap_touches": touch_sum - len(union_students),
        "overlap_rate": rate(touch_sum - len(union_students), denominator),
        "multi_port_students": multi_port_students,
        "multi_port_student_rate": rate(multi_port_students, denominator),
        "union_efficiency": rate(len(union_students), touch_sum),
    })

replacement_months = {row["month"] for row in updated_rows}
rows = [row for row in existing.get("rows", []) if row["month"] not in replacement_months]
rows.extend(updated_rows)
rows.sort(key=lambda row: row["month"])
if not rows:
    raise RuntimeError("No report rows were generated")

baseline = next((row for row in rows if row["month"] == TARGET_BASELINE_MONTH), rows[-1])
target_total = 0.65
auxiliary_rate = baseline["rates"]["student"] + baseline["rates"]["other"]
major_rate = baseline["rates"]["CC"] + baseline["rates"]["SS"] + baseline["rates"]["LP"]
required_touch_sum = target_total / baseline["union_efficiency"]
exact = {
    port: required_touch_sum * baseline["rates"][port] / major_rate
    for port in ("CC", "SS", "LP")
}
recommended = {"CC": 0.33, "SS": 0.38, "LP": 0.25}
recommended_touch_sum = sum(recommended.values())

result = {
    "source_as_of": SOURCE_AS_OF,
    "range": {"start": rows[0]["month"], "end": rows[-1]["month"]},
    "ports": list(PORTS),
    "rows": rows,
    "target": {
        "total": target_total,
        "baseline_period": baseline["month"],
        "baseline": {**baseline["rates"], "total": baseline["binding_rate"], "touch_sum": baseline["touch_sum_rate"], "union_efficiency": baseline["union_efficiency"]},
        "exact": exact,
        "recommended": recommended,
        "auxiliary_rate": auxiliary_rate,
        "required_touch_sum": required_touch_sum,
        "recommended_touch_sum": recommended_touch_sum,
        "estimated_total": recommended_touch_sum * baseline["union_efficiency"],
    },
    "quality": {
        "valid_binding_rows": valid_binding_rows,
        "pay_students": len(pay_students),
        "latest_touch_sum": rows[-1]["touch_sum"],
        "latest_bound_students": rows[-1]["bound_students"],
        "latest_overlap_touches": rows[-1]["overlap_touches"],
    },
}

JSON_OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
JS_OUTPUT.write_text(
    "window.FIXED_PLAN_M01_ALL_DATA = "
    + json.dumps(result, ensure_ascii=False, separators=(",", ":"))
    + ";\n",
    encoding="utf-8",
)
print(json.dumps({"updated_months": sorted(replacement_months), "latest": rows[-1]}, ensure_ascii=False, indent=2))
