import json
import os
from collections import Counter, defaultdict
from datetime import date, datetime

from openpyxl import load_workbook


BIND_PATH = os.environ["FP_BIND"]
PAY_PATH = os.environ["FP_PAY"]
BASE_DIR = os.path.dirname(__file__)
JSON_OUTPUT = os.path.join(BASE_DIR, "fixed-plan-ticket-data.json")
JS_OUTPUT = os.path.join(BASE_DIR, "fixed-plan-ticket-data.js")
START_MONTH = "2025-02"
END_MONTH = "2026-07"
PORTS = ("CC", "SS", "LP", "student", "other")


def normalize_id(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def to_datetime(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, str) and value.strip():
        text = value.strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                pass
    return None


def month_key(value):
    parsed = to_datetime(value)
    return parsed.strftime("%Y-%m") if parsed else None


def month_index(value):
    year, month = map(int, value.split("-"))
    return year * 12 + month


def month_range(start, end):
    current = month_index(start)
    final = month_index(end)
    values = []
    while current <= final:
        year, offset = divmod(current - 1, 12)
        values.append(f"{year:04d}-{offset + 1:02d}")
        current += 1
    return values


def normalize_port(value):
    lowered = str(value or "").strip().lower()
    if lowered in ("cc", "ss", "lp"):
        return lowered.upper()
    if lowered == "student":
        return "student"
    return "other"


def rate(count, denominator):
    return count / denominator if denominator else 0


def load_pay_students():
    workbook = load_workbook(PAY_PATH, read_only=True, data_only=True)
    sheet = workbook.active
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    positions = {name: index for index, name in enumerate(headers)}
    students = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        student_id = normalize_id(row[positions["stdt_id"]])
        pay_date = to_datetime(row[positions["first_1v1_nml_pay_date"]])
        if not student_id or not pay_date:
            continue
        current = students.get(student_id)
        if current is None or pay_date < current:
            students[student_id] = pay_date
    return students


def load_daily_first_tickets():
    workbook = load_workbook(BIND_PATH, read_only=True, data_only=True)
    sheet = workbook.active
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    positions = {name: index for index, name in enumerate(headers)}
    tickets = {}
    valid_rows = 0
    same_day_duplicate_rows = 0
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        student_id = normalize_id(row[positions["b.student_id"]])
        bind_time = to_datetime(row[positions["固定绑定时间"]])
        if not student_id or not bind_time:
            continue
        valid_rows += 1
        key = (student_id, bind_time.date().isoformat())
        candidate = {
            "student_id": student_id,
            "time": bind_time,
            "port": normalize_port(row[positions["bind_operator_type"]]),
            "row_number": row_number,
        }
        current = tickets.get(key)
        if current is None:
            tickets[key] = candidate
        else:
            same_day_duplicate_rows += 1
            if (bind_time, row_number) < (current["time"], current["row_number"]):
                tickets[key] = candidate
    return list(tickets.values()), {
        "valid_binding_rows": valid_rows,
        "same_day_duplicate_rows_removed": same_day_duplicate_rows,
        "daily_unique_tickets": len(tickets),
    }


pay_students = load_pay_students()
tickets, quality = load_daily_first_tickets()
months = month_range(START_MONTH, END_MONTH)
cohort_students = defaultdict(list)
for student_id, pay_date in pay_students.items():
    cohort = month_key(pay_date)
    if START_MONTH <= cohort <= END_MONTH:
        cohort_students[cohort].append(student_id)

tickets_by_student = defaultdict(list)
missing_pay_tickets = 0
pre_pay_tickets = 0
post_cutoff_tickets = 0
for ticket in tickets:
    pay_date = pay_students.get(ticket["student_id"])
    if pay_date is None:
        missing_pay_tickets += 1
        continue
    cohort = month_key(pay_date)
    if not (START_MONTH <= cohort <= END_MONTH):
        continue
    binding_month = month_key(ticket["time"])
    delta = month_index(binding_month) - month_index(cohort)
    if delta < 0:
        pre_pay_tickets += 1
        continue
    if binding_month > END_MONTH:
        post_cutoff_tickets += 1
        continue
    tickets_by_student[ticket["student_id"]].append(ticket)

rows = []
for cohort in months:
    students = cohort_students.get(cohort, [])
    denominator = len(students)
    stage_counts = Counter()
    port_counts = {stage: Counter() for stage in ("M0", "M1", "M2plus")}
    students_with_ticket = set()
    for student_id in students:
        for ticket in tickets_by_student.get(student_id, []):
            delta = month_index(month_key(ticket["time"])) - month_index(cohort)
            stage = "M0" if delta == 0 else "M1" if delta == 1 else "M2plus"
            stage_counts[stage] += 1
            port_counts[stage][ticket["port"]] += 1
            students_with_ticket.add(student_id)
    stage_payload = {}
    for stage in ("M0", "M1", "M2plus"):
        stage_payload[stage] = {
            "count": stage_counts[stage],
            "rate": rate(stage_counts[stage], denominator),
            "ports": {
                port: {"count": port_counts[stage][port], "rate": rate(port_counts[stage][port], denominator)}
                for port in PORTS
            },
        }
    total = sum(stage_counts.values())
    rows.append({
        "cohort": cohort,
        "new_students": denominator,
        "students_with_ticket": len(students_with_ticket),
        "M0": stage_payload["M0"],
        "M1": stage_payload["M1"],
        "M2plus": stage_payload["M2plus"],
        "total": {"count": total, "rate": rate(total, denominator)},
        "maturity": {
            "M0": True,
            "M1": month_index(cohort) + 1 <= month_index(END_MONTH),
            "M2plus": month_index(cohort) + 2 <= month_index(END_MONTH),
        },
    })

mature_rows = [row for row in rows if row["maturity"]["M2plus"]]
mature_students = sum(row["new_students"] for row in mature_rows)
aggregate = {"new_students": mature_students}
for stage in ("M0", "M1", "M2plus"):
    count = sum(row[stage]["count"] for row in mature_rows)
    aggregate[stage] = {
        "count": count,
        "rate": rate(count, mature_students),
        "ports": {
            port: {
                "count": sum(row[stage]["ports"][port]["count"] for row in mature_rows),
                "rate": rate(sum(row[stage]["ports"][port]["count"] for row in mature_rows), mature_students),
            }
            for port in PORTS
        },
    }
aggregate_total = sum(aggregate[stage]["count"] for stage in ("M0", "M1", "M2plus"))
aggregate["total"] = {"count": aggregate_total, "rate": rate(aggregate_total, mature_students)}

result = {
    "source_as_of": END_MONTH,
    "range": {"start": START_MONTH, "end": END_MONTH},
    "ports": list(PORTS),
    "rows": rows,
    "aggregate": aggregate,
    "quality": {
        **quality,
        "pay_students": len(pay_students),
        "cohort_students": sum(row["new_students"] for row in rows),
        "pre_pay_daily_tickets_excluded": pre_pay_tickets,
        "post_cutoff_daily_tickets_excluded": post_cutoff_tickets,
        "daily_tickets_missing_from_pay_table": missing_pay_tickets,
    },
}

with open(JSON_OUTPUT, "w", encoding="utf-8") as handle:
    json.dump(result, handle, ensure_ascii=False, indent=2)
with open(JS_OUTPUT, "w", encoding="utf-8") as handle:
    handle.write("window.FIXED_PLAN_TICKET_DATA = ")
    json.dump(result, handle, ensure_ascii=False, separators=(",", ":"))
    handle.write(";\n")

print(json.dumps({"aggregate": aggregate, "quality": result["quality"], "rows": rows}, ensure_ascii=False, indent=2))
